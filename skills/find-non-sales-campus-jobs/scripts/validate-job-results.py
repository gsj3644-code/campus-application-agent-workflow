#!/usr/bin/env python3
"""Validate non-sales campus job search exports and checkpoint JSON."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path


REQUIRED_SHEETS = ["使用说明", "岗位总表", "本轮新增与变更", "官网监控池", "不建议投与失效", "搜索覆盖与日志"]
REQUIRED_COLUMNS = [
    "建议等级", "信息粒度", "公司", "岗位全名", "城市", "招聘届别", "核心职责", "硬性要求",
    "匹配原因", "现实风险", "当前状态", "截止日期", "核验日期", "官方链接", "站内搜索词", "来源层级",
    "发现方式", "初始渠道", "原始线索", "漏检原因", "学习规则ID",
]
SOURCE_QUALITY_COLUMNS = ["发现入口类型", "公众号身份", "回源状态", "回源链接"]
LEVELS = {"主投", "可以投", "冲刺", "不建议投"}
GRANULARITIES = {"精确岗位", "精确方向", "批次入口", "监控线索", "已失效岗位"}
STATUSES = {"官网明确日期", "官方公告明确日期", "招满即止", "未披露", "待复核", "已失效", "技术性待复核"}
SALES_TERMS = (
    "客户开发", "客户维护", "渠道拓展", "商务拓展", "招商", "拉新", "成交", "回款", "佣金",
    "收入指标", "销售额", "业绩指标", "产品营销", "市场开拓", "商机挖掘", "销售协同", "售前",
)
DISCOVERY_METHODS = {"AI搜索", "用户提供", "历史复查", "监控触发"}
MISS_REASONS = {
    "outside_scope", "company_pool_gap", "alias_or_entity_gap", "role_keyword_gap", "source_channel_gap",
    "ats_or_dynamic_gap", "time_window_gap", "ranking_or_filter_gap", "interrupted_or_technical",
    "not_a_miss", "unknown",
}
DISCOVERY_ENTRY_TYPES = {
    "公司官方公众号", "公司招聘公众号", "政府国资人社公众号", "高校就业公众号", "非官方招聘公众号",
    "搜狗微信搜索", "公司官网", "官方ATS", "政府公共平台", "高校就业网", "聚合平台", "其他",
}
WECHAT_IDENTITIES = {
    "公司官方", "公司招聘官方", "政府/国资/人社官方", "高校官方", "可信转载", "非官方聚合", "无法确认", "不适用",
}
BACKTRACE_STATUSES = {
    "已回源到职位详情", "已回源到官方ATS/招聘首页", "官方公众号原文完整", "仅官方公众号线索",
    "仅转载线索", "技术性无法回源", "不适用",
}
PROMOTION_BACKTRACE = {"已回源到职位详情", "已回源到官方ATS/招聘首页", "官方公众号原文完整"}
DIRECT_OFFICIAL_ENTRIES = {"公司官网", "官方ATS", "政府公共平台", "高校就业网"}
SOURCE_METRIC_COUNTS = {
    "wechat_discovered_leads", "wechat_unique_leads", "official_source_verified", "official_account_complete",
    "unverified_leads", "stale_or_duplicate_leads",
}


def norm(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_json(path: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(data, list):
        return data, {"_format": "json", "_checkpoint_present": False}
    if not isinstance(data, dict):
        raise ValueError("JSON 顶层必须是对象或岗位数组")
    jobs = data.get("jobs", data.get("岗位总表", []))
    if not isinstance(jobs, list):
        raise ValueError("jobs/岗位总表 必须是数组")
    meta = dict(data)
    meta["_format"] = "json"
    meta["_checkpoint_present"] = "coverage" in data
    return jobs, meta


def load_csv(path: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle)), {}


def load_xlsx(path: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("验证 XLSX 需要安装 openpyxl；可改用 JSON/CSV") from exc
    workbook = load_workbook(path, data_only=False)
    values_workbook = load_workbook(path, data_only=True)
    missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    if missing:
        raise ValueError("缺少工作表: " + "、".join(missing))
    sheet = workbook["岗位总表"]
    values_sheet = values_workbook["岗位总表"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], {"workbook": workbook}
    headers = [norm(v) for v in rows[0]]
    jobs = [dict(zip(headers, row)) for row in rows[1:] if any(norm(v) for v in row)]
    workbook_issues: list[str] = []
    for name in workbook.sheetnames:
        if name not in REQUIRED_SHEETS and ("今日清单" in name or name.endswith("岗位表")):
            workbook_issues.append(f"发现按城市或范围拆分的非标准工作表: {name}")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.data_type == "f":
                cached = values_sheet[cell.coordinate].value
                if cached is None or (isinstance(cached, str) and cached.startswith("#")):
                    workbook_issues.append(f"公式未计算或存在错误: 岗位总表!{cell.coordinate}")
    for index, header in enumerate(headers, start=1):
        if header in REQUIRED_COLUMNS:
            letter = sheet.cell(row=1, column=index).column_letter
            width = sheet.column_dimensions[letter].width
            if width is not None and width < 8:
                workbook_issues.append(f"关键字段列宽过窄，可能显示截断: {header}({width})")
    return jobs, {"workbook_issues": workbook_issues}


def validate_jobs(jobs: list[dict[str, object]], max_sprint_ratio: float, schema_version: int = 1) -> list[str]:
    errors: list[str] = []
    if jobs:
        missing_columns = [c for c in REQUIRED_COLUMNS if c not in jobs[0]]
        if missing_columns:
            errors.append("岗位总表缺少字段: " + "、".join(missing_columns))
        if schema_version >= 2:
            missing_source_columns = [c for c in SOURCE_QUALITY_COLUMNS if c not in jobs[0]]
            if missing_source_columns:
                errors.append("schema_version 2 岗位总表缺少来源质量字段: " + "、".join(missing_source_columns))
    seen: set[tuple[str, str, str, str]] = set()
    recommended = 0
    sprint = 0
    for index, job in enumerate(jobs, start=2):
        level = norm(job.get("建议等级"))
        grain = norm(job.get("信息粒度"))
        status = norm(job.get("当前状态"))
        if level not in LEVELS:
            errors.append(f"第 {index} 行建议等级无效: {level or '<空>'}")
        if grain not in GRANULARITIES:
            errors.append(f"第 {index} 行信息粒度无效: {grain or '<空>'}")
        if status not in STATUSES:
            errors.append(f"第 {index} 行当前状态无效: {status or '<空>'}")
        key = tuple(norm(job.get(k)).casefold() for k in ("公司", "岗位全名", "城市", "招聘届别"))
        if all(key):
            if key in seen:
                errors.append(f"第 {index} 行岗位重复: {' / '.join(key)}")
            seen.add(key)
        if level in {"主投", "可以投", "冲刺"}:
            recommended += 1
        if level == "冲刺":
            sprint += 1
        if level in {"主投", "可以投"} and not (norm(job.get("官方链接")) or norm(job.get("站内搜索词"))):
            errors.append(f"第 {index} 行{level}岗位缺少官方链接或站内搜索词")
        if level in {"主投", "可以投", "冲刺"}:
            text = " ".join(norm(job.get(k)) for k in ("岗位全名", "核心职责", "硬性要求"))
            hits = [term for term in SALES_TERMS if term in text]
            if hits:
                errors.append(f"第 {index} 行推荐岗位含销售风险词，需人工复核: {','.join(hits)}")
        if grain == "批次入口" and norm(job.get("岗位全名")) and "批次" not in norm(job.get("岗位全名")) and "项目" not in norm(job.get("岗位全名")):
            errors.append(f"第 {index} 行批次入口疑似被表述为精确岗位")
        discovery_method = norm(job.get("发现方式"))
        if discovery_method and discovery_method not in DISCOVERY_METHODS:
            errors.append(f"第 {index} 行发现方式无效: {discovery_method}")
        if discovery_method == "用户提供":
            if not norm(job.get("初始渠道")):
                errors.append(f"第 {index} 行用户提供岗位缺少初始渠道")
            if not norm(job.get("原始线索")):
                errors.append(f"第 {index} 行用户提供岗位缺少原始线索")
            reason = norm(job.get("漏检原因"))
            if reason not in MISS_REASONS:
                errors.append(f"第 {index} 行用户提供岗位漏检原因无效: {reason or '<空>'}")
            if not norm(job.get("学习规则ID")):
                errors.append(f"第 {index} 行用户提供岗位缺少学习规则ID或“不适用”")
        if schema_version >= 2:
            entry_type = norm(job.get("发现入口类型"))
            identity = norm(job.get("公众号身份"))
            backtrace = norm(job.get("回源状态"))
            if entry_type not in DISCOVERY_ENTRY_TYPES:
                errors.append(f"第 {index} 行发现入口类型无效: {entry_type or '<空>'}")
            if identity not in WECHAT_IDENTITIES:
                errors.append(f"第 {index} 行公众号身份无效: {identity or '<空>'}")
            if backtrace not in BACKTRACE_STATUSES:
                errors.append(f"第 {index} 行回源状态无效: {backtrace or '<空>'}")
            if level in {"主投", "可以投"}:
                direct_official = entry_type in DIRECT_OFFICIAL_ENTRIES and backtrace == "不适用"
                if backtrace not in PROMOTION_BACKTRACE and not direct_official:
                    errors.append(f"第 {index} 行{level}岗位未通过官方来源晋级门: {backtrace or '<空>'}")
    if recommended and sprint / recommended > max_sprint_ratio + 1e-9:
        errors.append(f"冲刺岗位比例 {sprint / recommended:.1%} 超过上限 {max_sprint_ratio:.1%}")
    return errors


def validate_checkpoint(meta: dict[str, object]) -> list[str]:
    errors: list[str] = []
    errors.extend(str(item) for item in meta.get("workbook_issues", []))
    if meta.get("_format") == "json" and meta.get("_checkpoint_present") is not True:
        return errors + ["JSON 结果缺少搜索覆盖检查点"]
    if not meta or "coverage" not in meta:
        return errors

    mode = norm(meta.get("search_mode") or "standard").lower()
    if mode not in {"quick", "standard", "deep"}:
        errors.append(f"search_mode 无效: {mode or '<空>'}")
        mode = "standard"
    completion_claimed = meta.get("completion_claimed") is True
    try:
        schema_version = int(meta.get("schema_version", 1))
    except (TypeError, ValueError):
        errors.append("schema_version 必须是整数")
        schema_version = 1
    if meta.get("full_search_claimed") is True:
        completion_claimed = True
        if mode != "deep":
            errors.append("full_search_claimed 只能用于 deep 档")

    coverage = meta.get("coverage")
    if not isinstance(coverage, list):
        errors.append("coverage 必须是数组")
        coverage = []
    incomplete = [item for item in coverage if not isinstance(item, dict) or item.get("status") != "completed"]

    discovery = meta.get("discovery_passes", [])
    if not isinstance(discovery, list):
        errors.append("discovery_passes 必须是数组")
        discovery = []
    completed_channels = {
        norm(item.get("channel"))
        for item in discovery
        if isinstance(item, dict) and item.get("status") == "completed" and norm(item.get("channel"))
    }
    companies = meta.get("company_pool", [])
    if not isinstance(companies, list):
        errors.append("company_pool 必须是数组")
        companies = []
    standard_company_incomplete = [
        item for item in companies
        if not isinstance(item, dict)
        or item.get("status") != "checked"
        or item.get("recruitment_channels_checked") is not True
    ]
    deep_company_incomplete = [
        item for item in companies
        if not isinstance(item, dict)
        or item.get("status") != "checked"
        or item.get("entity_expansion_checked") is not True
        or item.get("recruitment_channels_checked") is not True
    ]

    dynamic = meta.get("dynamic_sites", [])
    dynamic_incomplete = []
    if isinstance(dynamic, list):
        for site in dynamic:
            if (
                not isinstance(site, dict)
                or site.get("status") != "completed"
                or not all(site.get(k) is True for k in ("pagination_checked", "filters_checked", "keywords_checked", "details_checked"))
                or (
                    isinstance(site.get("expected_count"), int)
                    and isinstance(site.get("retrieved_count"), int)
                    and site.get("expected_count") != site.get("retrieved_count")
                )
            ):
                dynamic_incomplete.append(site)
    else:
        errors.append("dynamic_sites 必须是数组")

    wechat = meta.get("wechat_searches", [])
    wechat_incomplete = []
    if isinstance(wechat, list):
        for index, search in enumerate(wechat, start=1):
            if not isinstance(search, dict):
                errors.append(f"wechat_searches 第 {index} 项必须是对象")
                wechat_incomplete.append(search)
                continue
            if not norm(search.get("query")):
                errors.append(f"wechat_searches 第 {index} 项缺少 query")
            search_url = norm(search.get("search_url"))
            if "weixin.sogou.com/weixin?type=2" not in search_url:
                errors.append(f"wechat_searches 第 {index} 项不是搜狗微信文章搜索 URL")
            if not isinstance(search.get("pages_checked"), int) or search.get("pages_checked", 0) < 1:
                errors.append(f"wechat_searches 第 {index} 项 pages_checked 必须至少为 1")
            if not isinstance(search.get("result_count"), int) or search.get("result_count", -1) < 0:
                errors.append(f"wechat_searches 第 {index} 项 result_count 必须是非负整数")
            if not isinstance(search.get("article_urls", []), list):
                errors.append(f"wechat_searches 第 {index} 项 article_urls 必须是数组")
            if norm(search.get("status")) not in {"completed", "technical_review"}:
                errors.append(f"wechat_searches 第 {index} 项状态无效")
                wechat_incomplete.append(search)
            elif norm(search.get("status")) != "completed":
                wechat_incomplete.append(search)
    else:
        errors.append("wechat_searches 必须是数组")
        wechat = []

    blind_spots = meta.get("residual_blind_spots", [])
    if not isinstance(blind_spots, list):
        errors.append("residual_blind_spots 必须是数组")
        blind_spots = ["invalid"]

    metrics = meta.get("source_quality_metrics")
    if schema_version >= 2:
        if not isinstance(metrics, dict):
            errors.append("schema_version 2 缺少 source_quality_metrics")
            metrics = {}
        for key in SOURCE_METRIC_COUNTS:
            value = metrics.get(key)
            if not isinstance(value, int) or isinstance(value, bool) or value < 0:
                errors.append(f"source_quality_metrics.{key} 必须是非负整数")
        unique = metrics.get("wechat_unique_leads")
        verified = metrics.get("official_source_verified")
        unverified = metrics.get("unverified_leads")
        rate = metrics.get("official_verification_rate")
        if isinstance(unique, int) and isinstance(verified, int) and verified > unique:
            errors.append("official_source_verified 不能超过 wechat_unique_leads")
        if isinstance(unique, int) and isinstance(unverified, int) and unverified > unique:
            errors.append("unverified_leads 不能超过 wechat_unique_leads")
        if unique == 0:
            if rate is not None:
                errors.append("wechat_unique_leads 为 0 时 official_verification_rate 应为 null")
        elif isinstance(unique, int) and unique > 0 and isinstance(verified, int):
            if not isinstance(rate, (int, float)) or isinstance(rate, bool) or not 0 <= rate <= 1:
                errors.append("official_verification_rate 必须在 0 到 1 之间")
            elif abs(rate - verified / unique) > 0.001:
                errors.append("official_verification_rate 与 official_source_verified / wechat_unique_leads 不一致")

    rules = meta.get("learning_rules", [])
    rule_ids: set[str] = set()
    if not isinstance(rules, list):
        errors.append("learning_rules 必须是数组")
        rules = []
    for index, rule in enumerate(rules, start=1):
        if not isinstance(rule, dict):
            errors.append(f"learning_rules 第 {index} 项必须是对象")
            continue
        rule_id = norm(rule.get("id"))
        if not rule_id:
            errors.append(f"learning_rules 第 {index} 项缺少 id")
        elif rule_id in rule_ids:
            errors.append(f"learning_rules ID 重复: {rule_id}")
        else:
            rule_ids.add(rule_id)
        if norm(rule.get("status")) not in {"candidate", "active", "retired"}:
            errors.append(f"learning_rules 第 {index} 项状态无效")
        if norm(rule.get("reason")) not in MISS_REASONS:
            errors.append(f"learning_rules 第 {index} 项漏检原因无效")
        for key in ("trigger", "action", "scope"):
            if not norm(rule.get(key)):
                errors.append(f"learning_rules 第 {index} 项缺少 {key}")

    for row_index, job in enumerate(meta.get("jobs", []), start=2):
        if isinstance(job, dict) and norm(job.get("发现方式")) == "用户提供":
            rule_id = norm(job.get("学习规则ID"))
            if rule_id not in {"不适用", "待生成"} and rule_id not in rule_ids:
                errors.append(f"第 {row_index} 行引用的学习规则不存在: {rule_id}")

    if completion_claimed:
        if meta.get("scope_declared") is not True or not coverage or incomplete:
            errors.append("声明完成但覆盖范围未全部完成")
        if mode == "standard":
            if not wechat:
                errors.append("标准档缺少高优先级搜狗微信文章搜索记录")
            if wechat_incomplete and not blind_spots:
                errors.append("标准档的搜狗微信技术待办未记录到渠道盲区")
            if len(completed_channels) < 1:
                errors.append("标准档至少需要一个已完成的公司发现渠道")
            if not companies:
                errors.append("标准档缺少有限公司池")
            if standard_company_incomplete:
                errors.append(f"标准档公司池仍有 {len(standard_company_incomplete)} 项未完成招聘渠道检查")
            if meta.get("old_positions_reviewed") is not True or meta.get("monitor_pool_checked") is not True:
                errors.append("标准档尚未完成旧岗位或监控池复查")
            if dynamic_incomplete and not blind_spots:
                errors.append("标准档的动态网站待办未记录到渠道盲区")
        elif mode == "deep":
            if not wechat:
                errors.append("深度档缺少搜狗微信组合查询记录")
            elif wechat_incomplete:
                errors.append(f"深度档仍有 {len(wechat_incomplete)} 组搜狗微信查询未完成")
            if len(completed_channels) < 2:
                errors.append("深度档少于两个独立且已完成的公司发现渠道")
            if not companies:
                errors.append("深度档缺少公司级覆盖分母")
            if deep_company_incomplete:
                errors.append(f"深度档公司池仍有 {len(deep_company_incomplete)} 项未完成实体或招聘渠道检查")
            if meta.get("old_positions_reviewed") is not True or meta.get("monitor_pool_checked") is not True:
                errors.append("深度档尚未完成旧岗位或监控池复查")
            if dynamic_incomplete:
                errors.append(f"深度档仍有 {len(dynamic_incomplete)} 个动态网站未完成")
            if blind_spots:
                errors.append(f"深度档仍有 {len(blind_spots)} 个未处理渠道盲区")
            if schema_version >= 2 and isinstance(metrics, dict) and metrics.get("unverified_leads", 0) != 0:
                errors.append("深度档仍有未回源公众号独有线索")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path)
    parser.add_argument("--max-sprint-ratio", type=float, default=0.20)
    args = parser.parse_args()
    if not 0 <= args.max_sprint_ratio <= 1:
        parser.error("--max-sprint-ratio 必须在 0 到 1 之间")
    loaders = {".json": load_json, ".csv": load_csv, ".xlsx": load_xlsx}
    loader = loaders.get(args.path.suffix.lower())
    if loader is None:
        parser.error("仅支持 .json、.csv 或 .xlsx")
    try:
        jobs, meta = loader(args.path)
        try:
            schema_version = int(meta.get("schema_version", 1))
        except (TypeError, ValueError):
            schema_version = 1
        errors = validate_jobs(jobs, args.max_sprint_ratio, schema_version) + validate_checkpoint(meta)
    except (OSError, ValueError, RuntimeError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    if errors:
        for error in errors:
            print(f"ERROR: {error}")
        print(f"FAILED: {len(errors)} issue(s)")
        return 1
    print(f"PASS: validated {len(jobs)} job row(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
